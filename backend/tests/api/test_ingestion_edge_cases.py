import io
import os
import zipfile
import pytest
try:
    import pymupdf as fitz
except ImportError:
    import fitz
import docx
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def create_in_memory_docx() -> bytes:
    """Creates a real, valid in-memory DOCX document for test ingestion."""
    doc = docx.Document()
    doc.add_heading("RESIDENTIAL LEASE AGREEMENT", level=1)
    doc.add_paragraph("This Residential Lease Agreement is entered into on January 1, 2026.")
    doc.add_heading("Section 1: Rent and Due Date", level=2)
    doc.add_paragraph("Tenant agrees to pay monthly rent of $2,400.00 due on the first day of each month.")
    doc.add_heading("Section 2: Early Termination Penalty", level=2)
    doc.add_paragraph("In the event of early termination, tenant shall pay a liquidated damages penalty of three (3) months' rent.")
    
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def create_encrypted_pdf() -> bytes:
    """Creates a real, password-protected PDF in memory."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Confidential Lease Agreement")
    perm = fitz.PDF_PERM_ACCESSIBILITY
    encrypt_meth = fitz.PDF_ENCRYPT_AES_256
    pdf_bytes = doc.tobytes(
        deflate=True,
        encryption=encrypt_meth,
        owner_pw="secret123",
        user_pw="secret123",
        permissions=perm,
    )
    doc.close()
    return pdf_bytes


def test_zero_byte_file_upload_rejected():
    """Empty files with 0 bytes must be rejected as unsupported or invalid format."""
    resp = client.post(
        "/api/documents",
        files={"file": ("empty.pdf", io.BytesIO(b""), "application/pdf")},
    )
    assert resp.status_code == 415
    assert "Unsupported or invalid file format" in resp.json()["detail"]


def test_mime_spoofing_text_file_as_pdf():
    """Text file disguised as .pdf must be rejected by magic byte validation."""
    fake_content = b"This is plain text pretending to be a legal lease contract."
    resp = client.post(
        "/api/documents",
        files={"file": ("fake.pdf", io.BytesIO(fake_content), "application/pdf")},
    )
    assert resp.status_code == 415


def test_mime_spoofing_exe_as_pdf():
    """Executable binary disguised as .pdf must be rejected."""
    exe_content = b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff\x00\x00"
    resp = client.post(
        "/api/documents",
        files={"file": ("malware.pdf", io.BytesIO(exe_content), "application/pdf")},
    )
    assert resp.status_code == 415


def test_mime_spoofing_non_docx_zip():
    """A generic zip file without WordprocessingML parts named .docx must be rejected."""
    zip_stream = io.BytesIO()
    with zipfile.ZipFile(zip_stream, "w") as zf:
        zf.writestr("malicious_script.sh", "echo 'pwned'")
    zip_bytes = zip_stream.getvalue()

    resp = client.post(
        "/api/documents",
        files={"file": ("test.docx", io.BytesIO(zip_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 415


def test_encrypted_pdf_rejection():
    """Encrypted/password-protected PDFs must be rejected safely without exposing stack traces."""
    enc_pdf = create_encrypted_pdf()
    resp = client.post(
        "/api/documents",
        files={"file": ("protected.pdf", io.BytesIO(enc_pdf), "application/pdf")},
    )
    assert resp.status_code == 500
    assert "Failed to process document" in resp.json()["detail"]


def test_corrupt_pdf_header_handling():
    """A file starting with %PDF- but containing corrupted body must fail gracefully."""
    corrupt_pdf = b"%PDF-1.4\n%corrupted bytes \x00\xff\xfe\xaa garbage that is not a valid pdf stream"
    resp = client.post(
        "/api/documents",
        files={"file": ("corrupt.pdf", io.BytesIO(corrupt_pdf), "application/pdf")},
    )
    assert resp.status_code == 500
    assert "Failed to process document" in resp.json()["detail"]


def test_path_traversal_filename_sanitization():
    """Filenames with directory traversal sequences must be sanitized to safe basenames."""
    docx_bytes = create_in_memory_docx()
    malicious_filename = "../../../../etc/passwd.docx"
    resp = client.post(
        "/api/documents",
        files={"file": (malicious_filename, io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 201
    doc_data = resp.json()
    # Stored filename must not contain directory traversal slashes
    assert "/" not in doc_data["filename"]
    assert "\\" not in doc_data["filename"]
    assert ".." not in doc_data["filename"]


def test_valid_docx_full_ingestion_and_chunking():
    """Valid DOCX upload must extract pages, chunk document, and index into database."""
    docx_bytes = create_in_memory_docx()
    resp = client.post(
        "/api/documents",
        files={"file": ("legitimate_lease.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 201
    doc_data = resp.json()
    assert doc_data["status"] == "ready"
    assert doc_data["page_count"] >= 1

    doc_id = doc_data["id"]
    # Verify chunks exist
    chunks_resp = client.get(f"/api/documents/{doc_id}/chunks")
    assert chunks_resp.status_code == 200
    chunks = chunks_resp.json()
    assert len(chunks) > 0
    assert any("Rent" in c["text"] or "$2,400" in c["text"] for c in chunks)
